"""Loader for Word documents (.docx), paragraph text only."""

from pathlib import Path

import docx

from rag_generator.domain import Document

PARAGRAPH_SEPARATOR = "\n\n"


class DocxLoader:
    def load(self, path: Path) -> Document:
        paragraphs = [p.text for p in docx.Document(str(path)).paragraphs if p.text.strip()]
        return Document(source=path.name, text=PARAGRAPH_SEPARATOR.join(paragraphs))
